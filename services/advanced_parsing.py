"""Advanced document parsing for PDF, DOCX, PPTX."""

from __future__ import annotations

from typing import Any


def extract_text_from_pdf(file_content: bytes) -> tuple[str, dict[str, Any]]:
    """Extract text from PDF using PyMuPDF."""
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:
        raise ImportError("PyMuPDF required: pip install pymupdf") from exc
    
    doc = fitz.open(stream=file_content, filetype="pdf")
    text_parts = []
    metadata = {"pages": len(doc), "format": "pdf"}
    
    for page_num, page in enumerate(doc):
        text_parts.append(page.get_text())
    
    doc.close()
    return "\n\n".join(text_parts), metadata


def extract_text_from_docx(file_content: bytes) -> tuple[str, dict[str, Any]]:
    """Extract text from DOCX using python-docx."""
    try:
        import io

        from docx import Document
    except ImportError as exc:
        raise ImportError("python-docx required: pip install python-docx") from exc
    
    doc = Document(io.BytesIO(file_content))
    text_parts = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    
    metadata = {"paragraphs": len(doc.paragraphs), "format": "docx"}
    return "\n\n".join(text_parts), metadata


def extract_text_from_pptx(file_content: bytes) -> tuple[str, dict[str, Any]]:
    """Extract text from PPTX using python-pptx."""
    try:
        import io

        from pptx import Presentation
    except ImportError as exc:
        raise ImportError("python-pptx required: pip install python-pptx") from exc
    
    prs = Presentation(io.BytesIO(file_content))
    text_parts = []
    
    for slide_num, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                text_parts.append(shape.text)
    
    metadata = {"slides": len(prs.slides), "format": "pptx"}
    return "\n\n".join(text_parts), metadata


def extract_text_auto(file_content: bytes, mime_type: str, filename: str | None = None) -> tuple[str, dict[str, Any]]:
    """Auto-detect format and extract text."""
    mime = mime_type.lower()
    ext = filename.lower().split('.')[-1] if filename and '.' in filename else ""
    
    if "pdf" in mime or ext == "pdf":
        return extract_text_from_pdf(file_content)
    elif "word" in mime or "docx" in mime or ext in ("docx", "doc"):
        if ext == "doc":
            # Basic fallback for .doc - try as text
            try:
                return file_content.decode('utf-8', errors='ignore'), {"format": "doc_fallback"}
            except Exception:  # noqa: BLE001
                raise ValueError("Old .doc format not fully supported, please convert to .docx")
        return extract_text_from_docx(file_content)
    elif "powerpoint" in mime or "pptx" in mime or ext in ("pptx", "ppt"):
        if ext == "ppt":
            try:
                return file_content.decode('utf-8', errors='ignore'), {"format": "ppt_fallback"}
            except Exception:  # noqa: BLE001
                raise ValueError("Old .ppt format not fully supported, please convert to .pptx")
        return extract_text_from_pptx(file_content)
    else:
        # Plain text
        try:
            return file_content.decode('utf-8'), {"format": "text"}
        except UnicodeError:
            return file_content.decode('latin-1'), {"format": "text_fallback"}
